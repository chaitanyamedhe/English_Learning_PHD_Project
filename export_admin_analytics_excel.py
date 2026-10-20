"""Export paper-ready admin analytics data to Excel.

Run with:
    python export_admin_analytics_excel.py

The script reads the live Django database and writes a workbook with:
    - Dashboard summary metrics
    - Student profile status
    - Learning progress by level
    - Assessment attempts
    - Learner analytics summary
    - Learning day analytics
    - Level transition analytics
    - Data gap / mismatch report
"""

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "english_learning.settings")

import django

django.setup()

from django.contrib.auth.models import User
from django.db.models import Avg, Count, Max, Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from home_page.models import StudentProfile
from learning.models import (
    LearnerAnalyticsSummary,
    LearningDayAnalytics,
    LevelTransitionAnalytics,
    UserAssessmentAttempt,
    UserCertificate,
    UserLearningProgress,
)


OUTPUT_PATH = Path(r"c:\English_learning\admin_analytics_export_paper_ready.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")


def normalize_value(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat(sep=" ")
        except TypeError:
            return value.isoformat()
    if isinstance(value, (list, dict)):
        return str(value)
    return value


def apply_table_style(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

def write_sheet(ws, headers, rows):
    ws.append(headers)
    widths = [len(str(header)) for header in headers]
    for row in rows:
        normalized_row = [normalize_value(item) for item in row]
        ws.append(normalized_row)
        for index, value in enumerate(normalized_row):
            widths[index] = max(widths[index], len(str(value)))
    apply_table_style(ws)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = min(width + 2, 40)


def level_label(raw_level: str) -> str:
    if not raw_level:
        return ""
    return raw_level.strip().lower().capitalize()


def main():
    wb = Workbook()
    wb.remove(wb.active)

    generated_at = timezone.now()

    total_students = User.objects.count()
    pretest_completed = StudentProfile.objects.filter(pretest_completed=True).count()
    certificates_issued = UserCertificate.objects.count()
    active_learners = LearnerAnalyticsSummary.objects.filter(completion_rate__gt=0).count()

    summary_stats = LearnerAnalyticsSummary.objects.aggregate(
        avg_accuracy=Avg("overall_accuracy"),
        avg_completion=Avg("completion_rate"),
        avg_engagement=Avg("engagement_score"),
        avg_speaking=Avg("speaking_confidence_score"),
        avg_improvement=Avg("overall_improvement_percentage"),
        avg_dropout_risk=Avg("dropout_risk_score"),
        total_learning_time=Sum("total_learning_time_seconds"),
        max_completion=Max("completion_rate"),
    )

    major_improvement_count = LearnerAnalyticsSummary.objects.filter(
        improvement_category="major_improvement"
    ).count()
    dropout_count = LearnerAnalyticsSummary.objects.filter(
        improvement_category="dropout"
    ).count()

    dashboard_rows = [
        ("generated_at", generated_at),
        ("total_students", total_students),
        ("pretest_completed", pretest_completed),
        ("active_learners", active_learners),
        ("certificates_issued", certificates_issued),
        ("avg_accuracy", round(summary_stats["avg_accuracy"] or 0, 2)),
        ("avg_completion", round(summary_stats["avg_completion"] or 0, 2)),
        ("avg_engagement", round(summary_stats["avg_engagement"] or 0, 2)),
        ("avg_speaking", round(summary_stats["avg_speaking"] or 0, 2)),
        ("avg_improvement", round(summary_stats["avg_improvement"] or 0, 2)),
        ("avg_dropout_risk", round(summary_stats["avg_dropout_risk"] or 0, 2)),
        ("major_improvement_count", major_improvement_count),
        ("dropout_count", dropout_count),
        ("total_learning_time_seconds", int(summary_stats["total_learning_time"] or 0)),
        ("max_completion_rate", round(summary_stats["max_completion"] or 0, 2)),
    ]

    ws = wb.create_sheet("Dashboard_Summary")
    write_sheet(ws, ["metric", "value"], dashboard_rows)
    ws["A1"].fill = TITLE_FILL
    ws["B1"].fill = TITLE_FILL

    ws = wb.create_sheet("StudentProfile")
    profile_headers = [
        "user_id",
        "username",
        "email",
        "institute",
        "department",
        "year",
        "listening_completed",
        "speaking_completed",
        "reading_completed",
        "writing_completed",
        "pretest_completed",
        "pretest_completed_at",
        "current_profile_level",
    ]
    profile_rows = []
    for profile in StudentProfile.objects.select_related("user").order_by("user_id"):
        profile_rows.append(
            (
                profile.user_id,
                profile.user.username,
                profile.user.email,
                profile.institute,
                profile.department,
                profile.year,
                profile.listening_completed,
                profile.speaking_completed,
                profile.reading_completed,
                profile.writing_completed,
                profile.pretest_completed,
                profile.pretest_completed_at,
                level_label(profile.level),
            )
        )
    write_sheet(ws, profile_headers, profile_rows)

    ws = wb.create_sheet("UserLearningProgress")
    progress_headers = [
        "user_id",
        "username",
        "level",
        "completed_days_count",
        "completed_days",
        "current_day",
        "streak_days",
        "last_completed_date",
        "started_at",
        "completed_at",
        "completion_percentage",
        "certificate_issued",
        "toolkit_completed",
        "toolkit_score",
        "toolkit_total_tools_completed",
        "has_seen_intro",
        "updated_at",
    ]
    progress_rows = []
    for progress in UserLearningProgress.objects.select_related("user").order_by("user_id", "level"):
        progress_rows.append(
            (
                progress.user_id,
                progress.user.username,
                progress.level,
                len(progress.completed_days or []),
                progress.completed_days,
                progress.current_day,
                progress.streak_days,
                progress.last_completed_date,
                progress.started_at,
                progress.completed_at,
                round(progress.completion_percentage(), 2),
                progress.certificate_issued,
                progress.toolkit_completed,
                progress.toolkit_score,
                progress.toolkit_total_tools_completed,
                progress.has_seen_intro,
                progress.updated_at,
            )
        )
    write_sheet(ws, progress_headers, progress_rows)

    ws = wb.create_sheet("UserAssessmentAttempt")
    attempt_headers = [
        "user_id",
        "username",
        "assessment_level",
        "score",
        "total_questions",
        "percentage",
        "passed",
        "started_at",
        "completed_at",
        "time_taken_seconds",
    ]
    attempt_rows = []
    for attempt in UserAssessmentAttempt.objects.select_related("user", "assessment").order_by("user_id", "-completed_at"):
        attempt_rows.append(
            (
                attempt.user_id,
                attempt.user.username,
                attempt.get_level_display(),
                attempt.score,
                attempt.total_questions,
                round(attempt.percentage, 2),
                attempt.passed,
                attempt.started_at,
                attempt.completed_at,
                attempt.time_taken_seconds,
            )
        )
    write_sheet(ws, attempt_headers, attempt_rows)

    ws = wb.create_sheet("LearnerAnalyticsSummary")
    summary_headers = [
        "user_id",
        "username",
        "email",
        "total_learning_time_seconds",
        "completed_days",
        "completion_rate",
        "overall_accuracy",
        "listening_score",
        "speaking_confidence_score",
        "reading_accuracy_score",
        "writing_improvement_score",
        "grammar_score",
        "vocabulary_engagement_score",
        "game_performance_score",
        "beginner_assessment_score",
        "intermediate_assessment_score",
        "advanced_toolkit_score",
        "pretest_score",
        "posttest_score",
        "overall_improvement_percentage",
        "engagement_score",
        "consistency_score",
        "retention_score",
        "dropout_risk_score",
        "total_retry_count",
        "strongest_section",
        "weakest_section",
        "improvement_category",
        "behavior_type",
        "certificate_issued",
        "last_updated",
    ]
    summary_rows = []
    for summary in LearnerAnalyticsSummary.objects.select_related("user").order_by("user_id"):
        summary_rows.append(
            (
                summary.user_id,
                summary.user.username,
                summary.user.email,
                summary.total_learning_time_seconds,
                summary.completed_days,
                round(summary.completion_rate, 2),
                round(summary.overall_accuracy, 2),
                round(summary.listening_score, 2),
                round(summary.speaking_confidence_score, 2),
                round(summary.reading_accuracy_score, 2),
                round(summary.writing_improvement_score, 2),
                round(summary.grammar_score, 2),
                round(summary.vocabulary_engagement_score, 2),
                round(summary.game_performance_score, 2),
                round(summary.beginner_assessment_score, 2),
                round(summary.intermediate_assessment_score, 2),
                round(summary.advanced_toolkit_score, 2),
                round(summary.pretest_score, 2),
                round(summary.posttest_score, 2),
                round(summary.overall_improvement_percentage, 2),
                round(summary.engagement_score, 2),
                round(summary.consistency_score, 2),
                round(summary.retention_score, 2),
                round(summary.dropout_risk_score, 2),
                summary.total_retry_count,
                summary.strongest_section,
                summary.weakest_section,
                summary.improvement_category,
                summary.behavior_type,
                summary.certificate_issued,
                summary.last_updated,
            )
        )
    write_sheet(ws, summary_headers, summary_rows)

    ws = wb.create_sheet("LearningDayAnalytics")
    day_headers = [
        "user_id",
        "username",
        "email",
        "level",
        "day_number",
        "stage",
        "started_at",
        "completed_at",
        "total_time_spent_seconds",
        "total_sections",
        "sections_completed",
        "total_attempts",
        "total_correct_attempts",
        "total_wrong_attempts",
        "overall_accuracy",
        "activity_data",
    ]
    day_rows = []
    for day in LearningDayAnalytics.objects.select_related("user").order_by("user_id", "level", "day_number", "stage"):
        day_rows.append(
            (
                day.user_id,
                day.user.username,
                day.user.email,
                day.level,
                day.day_number,
                day.stage,
                day.started_at,
                day.completed_at,
                day.total_time_spent_seconds,
                day.total_sections,
                day.sections_completed,
                day.total_attempts,
                day.total_correct_attempts,
                day.total_wrong_attempts,
                round(day.overall_accuracy, 2),
                day.activity_data,
            )
        )
    write_sheet(ws, day_headers, day_rows)

    ws = wb.create_sheet("LevelTransitionAnalytics")
    transition_headers = [
        "user_id",
        "username",
        "email",
        "initial_level",
        "final_level",
        "improvement_type",
        "strongest_section",
        "weakest_section",
        "pretest_score",
        "beginner_assessment_score",
        "intermediate_assessment_score",
        "advanced_toolkit_score",
        "posttest_score",
        "improvement_percentage",
        "transition_speed",
        "certificate_issued",
        "created_at",
        "updated_at",
    ]
    transition_rows = []
    for transition in LevelTransitionAnalytics.objects.select_related("user").order_by("user_id"):
        transition_rows.append(
            (
                transition.user_id,
                transition.user.username,
                transition.user.email,
                transition.initial_level,
                transition.final_level,
                transition.improvement_type,
                transition.strongest_section,
                transition.weakest_section,
                round(transition.pretest_score, 2),
                round(transition.beginner_assessment_score, 2),
                round(transition.intermediate_assessment_score, 2),
                round(transition.advanced_toolkit_score, 2),
                round(transition.posttest_score, 2),
                round(transition.improvement_percentage, 2),
                transition.transition_speed,
                transition.certificate_issued,
                transition.created_at,
                transition.updated_at,
            )
        )
    write_sheet(ws, transition_headers, transition_rows)

    ws = wb.create_sheet("UserCertificate")
    certificate_headers = [
        "user_id",
        "username",
        "email",
        "level",
        "certificate_code",
        "issued_at",
        "downloaded",
        "shared",
        "download_count",
        "share_count",
    ]
    certificate_rows = []
    for certificate in UserCertificate.objects.select_related("user").order_by("user_id", "level"):
        certificate_rows.append(
            (
                certificate.user_id,
                certificate.user.username,
                certificate.user.email,
                certificate.level,
                certificate.certificate_code,
                certificate.issued_at,
                certificate.downloaded,
                certificate.shared,
                certificate.download_count,
                certificate.share_count,
            )
        )
    write_sheet(ws, certificate_headers, certificate_rows)

    pretest_users = set(StudentProfile.objects.filter(pretest_completed=True).values_list("user_id", flat=True))
    progress_users = set(UserLearningProgress.objects.values_list("user_id", flat=True).distinct())
    summary_users = set(LearnerAnalyticsSummary.objects.values_list("user_id", flat=True).distinct())

    pretest_without_progress = sorted(pretest_users - progress_users)
    progress_without_pretest = sorted(progress_users - pretest_users)
    progress_without_summary = sorted(progress_users - summary_users)

    ws = wb.create_sheet("Data_Gaps")
    gap_headers = ["category", "user_id", "username", "email", "detail"]
    gap_rows = []

    for user_id in pretest_without_progress:
        user = User.objects.filter(id=user_id).first()
        gap_rows.append(
            (
                "pretest_completed_but_no_learning_progress",
                user_id,
                user.username if user else "",
                user.email if user else "",
                "Completed pretest but no UserLearningProgress row exists.",
            )
        )

    for user_id in progress_without_pretest:
        user = User.objects.filter(id=user_id).first()
        gap_rows.append(
            (
                "learning_progress_without_pretest_flag",
                user_id,
                user.username if user else "",
                user.email if user else "",
                "Has learning progress, but StudentProfile.pretest_completed is false.",
            )
        )

    for user_id in progress_without_summary:
        user = User.objects.filter(id=user_id).first()
        gap_rows.append(
            (
                "learning_progress_without_analytics_summary",
                user_id,
                user.username if user else "",
                user.email if user else "",
                "Has learning progress, but no LearnerAnalyticsSummary row exists.",
            )
        )

    write_sheet(ws, gap_headers, gap_rows)

    wb.save(OUTPUT_PATH)
    print(f"Workbook written to: {OUTPUT_PATH}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
    